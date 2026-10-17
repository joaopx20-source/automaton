#!/usr/bin/env python3
"""Agenda de uma campanha só para o YouTube, de trás para frente (último vídeo → 1),
4 vídeos por dia nos horários de sempre.

Uso:  python3 promo/kavix/youtube_reverso.py [campanha] [AAAA-MM-DD]
      (padrão: campanha 2, começando amanhã)
Saída: campanhaN/youtube_reverso/ (planilha, JSON, CSV e mensagem para o Codex)
"""
import csv
import datetime as dt
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
CAMP = "2"
SRC = OUT = BASE = None
HORARIOS = ["09:00", "12:30", "17:00", "20:30"]
POR_DIA = 4


def configurar(camp):
    global CAMP, SRC, OUT, BASE
    CAMP = camp
    pasta = "campanha" if camp == "1" else f"campanha{camp}"
    SRC = os.path.join(HERE, pasta, "plano_postagem.json")
    OUT = os.path.join(HERE, pasta, "youtube_reverso")
    BASE = ("https://github.com/joaopx20-source/automaton/raw/"
            f"claude/kavix-promotional-videos-e6f1vd/promo/kavix/{pasta}/youtube_reverso/")


def gerar(inicio):
    posts = sorted(json.load(open(SRC)), key=lambda p: -p["n"])
    linhas = []
    for k, p in enumerate(posts):
        dia_n = k // POR_DIA + 1
        dia = inicio + dt.timedelta(days=dia_n - 1)
        hora = HORARIOS[k % POR_DIA]
        linhas.append({
            "ordem": k + 1, "video_n": p["n"], "dia": dia_n,
            "data": dia.strftime("%d/%m/%Y"), "hora": hora,
            "data_hora": f"{dia.isoformat()}T{hora}:00-03:00",
            "arquivo": p["arquivo"], "url_download": p["url_download"],
            "titulo_youtube": p["titulo_youtube"],
            "descricao_youtube": p["descricao_youtube"],
            "legenda_redes": p["legenda_instagram_facebook"],
            "status": "pendente",
        })
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "plano_youtube_reverso.json"), "w") as fh:
        json.dump(linhas, fh, ensure_ascii=False, indent=2)
    with open(os.path.join(OUT, "plano_youtube_reverso.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, list(linhas[0]))
        w.writeheader()
        w.writerows(linhas)

    wb = Workbook()
    ws = wb.active
    ws.title = "YouTube"
    cab = ["Ordem", "Vídeo nº", "Dia", "Data", "Hora", "Data/hora (Brasília)", "Arquivo",
           "Link do vídeo", "Título YouTube", "Descrição YouTube", "Status"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B3FD9")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for L in linhas:
        ws.append([L["ordem"], L["video_n"], L["dia"], L["data"], L["hora"], L["data_hora"],
                   L["arquivo"], L["url_download"], L["titulo_youtube"],
                   L["descricao_youtube"], L["status"]])
    for i, wdt in enumerate([8, 9, 6, 12, 7, 27, 34, 40, 50, 60, 12], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
    claro = PatternFill("solid", fgColor="EEEAFE")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value % 2:
            for c in ws[r]:
                c.fill = claro
        ws.cell(r, 10).alignment = Alignment(wrap_text=False)
    dv = DataValidation(type="list", formula1='"pendente,agendado,publicado,erro"')
    ws.add_data_validation(dv)
    dv.add(f"K2:K{ws.max_row}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(os.path.join(OUT, "calendario_youtube_reverso.xlsx"))

    fim = inicio + dt.timedelta(days=linhas[-1]["dia"] - 1)
    msg = f"""Preciso que você agende no meu canal do YouTube (KAVIX ARCADE) os {len(linhas)} Shorts da
campanha {CAMP}, DE TRÁS PARA FRENTE: começa pelo vídeo {linhas[0]["video_n"]} e termina no vídeo 1.
Eu já estou logado no YouTube Studio no navegador. NÃO use o Buffer.

ARQUIVOS:
- Planilha: {BASE}calendario_youtube_reverso.xlsx
- Mesmo conteúdo em JSON: {BASE}plano_youtube_reverso.json
- Cada linha tem o link do vídeo (.mp4) na coluna "Link do vídeo".

AGENDA:
- 4 vídeos por dia, às 09:00, 12:30, 17:00 e 20:30, HORÁRIO DE BRASÍLIA (UTC−3).
- De {inicio.strftime('%d/%m/%Y')} até {fim.strftime('%d/%m/%Y')} ({linhas[-1]['dia']} dias).
- Ordem: linha 1 da planilha = vídeo {linhas[0]["video_n"]} (dia 1, 09:00) ... linha {len(linhas)} = vídeo 1.

COMO AGENDAR CADA LINHA (YouTube Studio):
1. Criar → Enviar vídeos → escolha o .mp4 da linha.
2. Título = coluna "Título YouTube"; Descrição = coluna "Descrição YouTube" (copie exatamente).
3. Público: "Não, não é conteúdo para crianças".
4. Visibilidade → Programar → data e hora da linha (fuso de Brasília).
5. Confira antes se o vídeo já não está programado (não duplicar).
6. Marque a coluna "Status" (agendado / erro).

LIMITE DO YOUTUBE: se aparecer aviso de limite diário de envios, pare, anote em qual
linha parou e continue no dia seguinte a partir dela (as datas já programadas não mudam).

No final, me mande um resumo: quantos foram agendados e quais deram erro.
"""
    with open(os.path.join(OUT, "MENSAGEM_PARA_CODEX.txt"), "w") as fh:
        fh.write(msg)
    print("ok:", len(linhas), "vídeos de", inicio, "a", fim)


if __name__ == "__main__":
    args = sys.argv[1:]
    configurar(args.pop(0) if args and args[0].isdigit() else "2")
    ini = (dt.date.fromisoformat(args[0]) if args
           else dt.date.today() + dt.timedelta(days=1))
    gerar(ini)
