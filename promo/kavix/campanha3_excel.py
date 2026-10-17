#!/usr/bin/env python3
"""Planilha e mensagem de agendamento da campanha 3 (300 reels, 4 por dia, 75 dias).

Uso:  python3 promo/kavix/campanha3_excel.py
Saída: campanha3/calendario_campanha3.xlsx e campanha3/MENSAGEM_PARA_CODEX.txt
(roda depois de `campanha3.py plano`)
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "campanha3")
BASE = ("https://github.com/joaopx20-source/automaton/raw/"
        "claude/kavix-promotional-videos-e6f1vd/promo/kavix/campanha3/")
ROXO = PatternFill("solid", fgColor="5B3FD9")
AMARELO = PatternFill("solid", fgColor="FFF2A8")


def cab(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ROXO
        c.alignment = Alignment(horizontal="center", vertical="center")


def main():
    posts = json.load(open(os.path.join(OUT, "plano_postagem.json")))
    n, dias = len(posts), posts[-1]["dia_campanha"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Inicio"
    linhas = [
        ["Campanha 3 KAVIX ARCADE: 300 reels (jogos, skins e ranked)", None, None],
        [None, None, None],
        ["Data de início (Dia 1):", None, "← digite a data aqui (ex.: 08/11/2026)"],
        [None, None, None],
        ["Posts por dia", "4, às 09:00, 12:30, 17:00 e 20:30 (horário de Brasília)", None],
        ["Total", f"{n} posts em {dias} dias", None],
        ["Redes", "Instagram (Reels), Facebook (Reels), TikTok, YouTube (Shorts): o mesmo vídeo "
                  "nas 4, no mesmo horário", None],
        ["Último dia", f'=IF(B3="","(defina a data de início)",B3+{dias - 1})', None],
        ["Vídeos", "Coluna url_download da aba Calendario (link direto do arquivo .mp4)", None],
        ["Agendar em", "Meta Business Suite (Instagram + Facebook), TikTok Studio, YouTube Studio "
                       "(sem Buffer)", None],
        ["Limites", "TikTok agenda até 10 dias à frente (fazer em lotes); Meta até 29 dias; "
                    "YouTube tem limite diário de envios", None],
    ]
    for L in linhas:
        ws.append(L)
    ws["A1"].font = Font(bold=True, size=14)
    ws["B3"].fill = AMARELO
    ws["B3"].number_format = "DD/MM/YYYY"
    ws["B8"].number_format = "DD/MM/YYYY"
    for r in range(3, 12):
        ws.cell(r, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 40

    cal = wb.create_sheet("Calendario")
    cols = ["n", "dia_campanha", "data", "hora", "data_hora", "formato", "assunto", "arquivo",
            "url_download", "legenda_instagram_facebook", "legenda_tiktok", "titulo_youtube",
            "descricao_youtube"]
    cab(cal, cols)
    for r, p in enumerate(posts, 2):
        cal.append([p["n"], p["dia_campanha"],
                    f'=IF(Inicio!$B$3="","",Inicio!$B$3+B{r}-1)', p["hora"],
                    f'=IF(C{r}="","",TEXT(C{r},"YYYY-MM-DD")&"T"&D{r}&":00-03:00")',
                    p["formato"], p["assunto"], p["arquivo"], p["url_download"],
                    p["legenda_instagram_facebook"], p["legenda_tiktok"], p["titulo_youtube"],
                    p["descricao_youtube"]])
        cal.cell(r, 3).number_format = "DD/MM/YYYY"
    for i, w in enumerate([6, 8, 12, 7, 26, 13, 22, 40, 50, 60, 60, 50, 60], 1):
        cal.column_dimensions[cal.cell(1, i).column_letter].width = w
    cal.freeze_panes = "C2"
    cal.auto_filter.ref = cal.dimensions

    ct = wb.create_sheet("Controle")
    cab(ct, ["n", "data", "hora", "Instagram", "Facebook", "TikTok", "YouTube", "observação"])
    for r in range(2, n + 2):
        ct.append([f"=Calendario!A{r}", f"=Calendario!C{r}", f"=Calendario!D{r}",
                   "pendente", "pendente", "pendente", "pendente", None])
        ct.cell(r, 2).number_format = "DD/MM/YYYY"
    dv = DataValidation(type="list", formula1='"pendente,agendado,publicado,erro"',
                        allow_blank=True)
    ct.add_data_validation(dv)
    dv.add(f"D2:G{n + 1}")
    resumo = [("Pendentes", f'=COUNTIF(D2:G{n + 1},"pendente")'),
              ("Agendados (total)", f'=COUNTIF(D2:G{n + 1},"agendado")+COUNTIF(D2:G{n + 1},"publicado")'),
              ("Erros", f'=COUNTIF(D2:G{n + 1},"erro")')]
    for k, (a, b) in enumerate(resumo, 2):
        ct.cell(k, 10, a).font = Font(bold=True)
        ct.cell(k, 11, b)
    ct.column_dimensions["J"].width = 20
    ct.freeze_panes = "A2"
    wb.save(os.path.join(OUT, "calendario_campanha3.xlsx"))

    msg = f"""Preciso que você agende a campanha 3 do KAVIX ARCADE: {n} reels nas minhas redes
(Instagram @kavixarcade, Facebook, TikTok e YouTube Shorts). Eu já estou logado em todas
no navegador. NÃO use o Buffer: agende nas ferramentas oficiais de cada rede.

ARQUIVOS (repositório público no GitHub):
- Planilha: {BASE}calendario_campanha3.xlsx
- Plano em JSON (mesmo conteúdo): {BASE}plano_postagem.json
- Vídeos: coluna "url_download" da planilha (um .mp4 por post).

DATA DE INÍCIO: [COLOQUE AQUI, ex.: 08/11/2026]
(Na planilha, digite essa data na célula amarela B3 da aba "Inicio". As colunas "data" e
"data_hora" da aba "Calendario" se calculam sozinhas. Dia N = início + N−1 dias.)

REGRAS:
1. 4 posts por dia às 09:00, 12:30, 17:00 e 20:30, HORÁRIO DE BRASÍLIA (UTC−3). São {dias} dias.
2. O mesmo vídeo vai para as 4 redes no mesmo horário.
3. Instagram + Facebook: Meta Business Suite → Planejador → Criar Reel (marque os dois),
   legenda = coluna "legenda_instagram_facebook". A Meta agenda até 29 dias à frente.
4. TikTok: TikTok Studio → Enviar → Agendar, legenda = "legenda_tiktok".
   O TikTok só agenda até 10 dias à frente: faça em lotes e volte para o resto depois.
5. YouTube: YouTube Studio → Enviar vídeo → título = "titulo_youtube",
   descrição = "descricao_youtube", "não é conteúdo para crianças", Programar = data_hora.
   Se aparecer limite diário de envios, pare e continue no dia seguinte.
6. Use exatamente o vídeo e o texto de cada linha; não mude legendas, hashtags nem a ordem.
7. Antes de agendar, confira se o post já não está agendado (não duplicar).
8. Marque o andamento na aba "Controle" (pendente / agendado / publicado / erro)
   e me mande um resumo no final: quantos agendados por rede e quais deram erro.
"""
    with open(os.path.join(OUT, "MENSAGEM_PARA_CODEX.txt"), "w") as fh:
        fh.write(msg)
    print("ok:", n, "posts,", dias, "dias")


if __name__ == "__main__":
    main()
